import os
import sqlite3
import json
import time
import numpy as np
import ollama
from pathlib import Path

# --- Configuration ---
WORKSPACE_DIR = Path(__file__).parent / "Alfred_Workspace"
LIBRARY_DIR = WORKSPACE_DIR / "Library"
DB_PATH = WORKSPACE_DIR / "library_index.sqlite"
EMBEDDING_MODEL = "nomic-embed-text"
LLM_MODEL = "qwen2.5-coder:3b"

# Ensure directories exist
os.makedirs(LIBRARY_DIR, exist_ok=True)

def _get_db():
    conn = sqlite3.connect(str(DB_PATH))
    conn.execute('''
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT UNIQUE,
            last_modified REAL
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_id INTEGER,
            chunk_index INTEGER,
            text_content TEXT,
            embedding BLOB,
            FOREIGN KEY(doc_id) REFERENCES documents(id) ON DELETE CASCADE
        )
    ''')
    # Enable foreign keys
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def _get_embedding(text: str) -> np.ndarray:
    try:
        import shared
        response = shared.local_client.embeddings(model=EMBEDDING_MODEL, prompt=text)
        return np.array(response["embedding"], dtype=np.float32)
    except Exception as e:
        print(f"[Scholar Error] Embedding failed: {e}")
        return np.zeros(768, dtype=np.float32) # nomic-embed-text uses 768 dims

def _chunk_text(text: str, chunk_size: int = 500, overlap: int = 100):
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
        i += chunk_size - overlap
    return chunks

def _extract_text_from_file(filepath: Path) -> str:
    ext = filepath.suffix.lower()
    text = ""
    try:
        if ext == '.txt':
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        elif ext == '.pdf':
            import PyPDF2
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
        elif ext in ['.doc', '.docx']:
            import docx
            doc = docx.Document(filepath)
            for para in doc.paragraphs:
                text += para.text + "\n"
    except Exception as e:
        print(f"[Scholar] Error reading {filepath.name}: {e}")
    
    return text.strip()

def sync_library():
    """
    Scans the Library folder for new or modified files.
    Extracts text, chunks it, and vectorizes it.
    """
    print("[Scholar] Syncing library...")
    conn = _get_db()
    c = conn.cursor()
    
    # Get tracked files
    c.execute("SELECT id, filename, last_modified FROM documents")
    tracked_docs = {row[1]: {'id': row[0], 'mtime': row[2]} for row in c.fetchall()}
    
    current_files = set()
    
    for item in LIBRARY_DIR.iterdir():
        if item.is_dir():
            continue
            
        ext = item.suffix.lower()
        if ext not in ['.txt', '.pdf', '.docx', '.doc']:
            continue
            
        filename = item.name
        current_files.add(filename)
        mtime = item.stat().st_mtime
        
        needs_update = False
        if filename not in tracked_docs:
            needs_update = True
        elif tracked_docs[filename]['mtime'] < mtime:
            needs_update = True
            # Delete old chunks
            c.execute("DELETE FROM documents WHERE filename = ?", (filename,))
            conn.commit()
            
        if needs_update:
            print(f"[Scholar] Ingesting document: {filename}")
            text = _extract_text_from_file(item)
            if not text:
                print(f"[Scholar] Warning: No text extracted from {filename}")
                continue
                
            c.execute("INSERT INTO documents (filename, last_modified) VALUES (?, ?)", (filename, mtime))
            doc_id = c.lastrowid
            
            chunks = _chunk_text(text)
            for idx, chunk in enumerate(chunks):
                emb = _get_embedding(chunk)
                c.execute("INSERT INTO chunks (doc_id, chunk_index, text_content, embedding) VALUES (?, ?, ?, ?)",
                          (doc_id, idx, chunk, emb.tobytes()))
            conn.commit()
            
    # Remove deleted files from DB (batch operation instead of N+1 loop)
    deleted_files = [f for f in tracked_docs if f not in current_files]
    if deleted_files:
        for f in deleted_files:
            print(f"[Scholar] Removing deleted document: {f}")
        placeholders = ",".join("?" * len(deleted_files))
        c.execute(f"DELETE FROM documents WHERE filename IN ({placeholders})", deleted_files)
        conn.commit()
    conn.close()
    print("[Scholar] Library sync complete.")

def query_library(query: str, top_k: int = 3) -> str:
    """
    Embeds the query, finds the most relevant document chunks via cosine similarity,
    and returns a summarized answer using the LLM.
    If no relevant information is found, returns a failure string so the agent can fallback to web search.
    """
    sync_library() # Always sync before querying to catch newly dropped files
    
    query_emb = _get_embedding(query)
    
    conn = _get_db()
    c = conn.cursor()
    
    c.execute("""
        SELECT c.text_content, c.embedding, d.filename 
        FROM chunks c 
        JOIN documents d ON c.doc_id = d.id
    """)
    rows = c.fetchall()
    conn.close()
    
    if not rows:
        return "The Library is completely empty. There are no documents to search."
        
    results = []
    for row in rows:
        text_content, emb_bytes, filename = row
        chunk_emb = np.frombuffer(emb_bytes, dtype=np.float32)
        
        # Cosine similarity
        norm_q = np.linalg.norm(query_emb)
        norm_c = np.linalg.norm(chunk_emb)
        if norm_q == 0 or norm_c == 0:
            sim = 0
        else:
            sim = np.dot(query_emb, chunk_emb) / (norm_q * norm_c)
            
        results.append({
            "text": text_content,
            "filename": filename,
            "similarity": float(sim)
        })
        
    # Sort by highest similarity
    results.sort(key=lambda x: x["similarity"], reverse=True)
    top_results = results[:top_k]
    
    # If the best match is very low similarity, the answer is likely not in the library
    if top_results[0]["similarity"] < 0.4:
        return "I searched your local Library documents, but could not find any relevant information. Please search the web instead."
        
    # Build context for the LLM
    context_str = "\n\n".join([f"[Source: {r['filename']}]\n{r['text']}" for r in top_results])
    
    prompt = f"""You are Alfred, a strict reading comprehension assistant.
Your ONLY job is to extract the answer to the Question from the 'Context from local Library' below.
RULES:
1. If the exact answer is not found in the Context text, you MUST reply with exactly: "NOT_FOUND"
2. Do NOT use outside knowledge.
3. Do NOT guess. 
4. Do NOT hallucinate personal details.

Question: {query}

Context from local Library:
{context_str}

Answer:"""

    try:
        import llm_engine
        res = llm_engine.chat(
            messages=[{'role': 'user', 'content': prompt}],
            options={'temperature': 0.1, 'num_predict': 200}
        )
        answer = res['message']['content'].strip()
        
        if "NOT_FOUND" in answer:
             return "I checked your local library documents but could not find the answer to that. Please search the web instead."
             
        return f"According to your local library files:\n{answer}"
        
    except Exception as e:
        return f"Error querying the library LLM: {e}"
